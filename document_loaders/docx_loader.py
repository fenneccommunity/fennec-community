# -*- coding: utf-8 -*-
"""
Microsoft Word Document Loader (.docx) 
"""

from typing import List, Optional
import logging
from .base_loader import BaseFileLoader, LoadedDocument
from .config_loader import DocxLoaderConfig

logger = logging.getLogger(__name__)


class DocxLoader(BaseFileLoader):
    """
    Load Microsoft Word (.docx) documents.
    """

    SUPPORTED_EXTENSIONS = [".docx", ".doc"]

    def __init__(self, file_path: str, config: Optional[DocxLoaderConfig] = None):
        super().__init__(file_path=file_path)
        self.config = config or DocxLoaderConfig()

    def load(self) -> List[LoadedDocument]:
        if self.file_path.suffix.lower() == ".doc":
            return self._load_legacy_doc()
        return self._load_docx()

    def _load_docx(self) -> List[LoadedDocument]:
        try:
            from docx import Document as DocxDocument
        except ImportError:
            raise ImportError("python-docx is required. Install: pip install python-docx")

        doc = DocxDocument(str(self.file_path))
        parts = []

        if self.config.include_headers:
            for section in doc.sections:
                text = self._extract_header_footer(section.header)
                if text:
                    parts.append(f"[Header]: {text}")

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            style = para.style.name if para.style else "Normal"
            if "Heading" in style:
                parts.append(f"\n{'#' * self._heading_level(style)} {text}\n")
            else:
                parts.append(text)

        if self.config.include_tables:
            for table in doc.tables:
                table_text = self._extract_table(table)
                if table_text:
                    parts.append("\n[Table]:\n" + table_text)

        if self.config.include_footers:
            for section in doc.sections:
                text = self._extract_header_footer(section.footer)
                if text:
                    parts.append(f"[Footer]: {text}")

        full_text = "\n".join(parts).strip()
        props = doc.core_properties
        meta = self._build_file_metadata(
            title=props.title or "",
            author=props.author or "",
            subject=props.subject or "",
            created=str(props.created or ""),
            modified=str(props.modified or ""),
            paragraph_count=len(doc.paragraphs),
            table_count=len(doc.tables),
        )
        return [LoadedDocument(page_content=full_text, metadata=meta)]

    def _extract_table(self, table) -> str:
        rows = []
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            rows.append(self.config.row_separator.join(cells))
        return self.config.table_separator.join(rows)

    def _extract_header_footer(self, hf) -> str:
        if hf is None:
            return ""
        texts = [p.text.strip() for p in hf.paragraphs if p.text.strip()]
        return " ".join(texts)

    def _heading_level(self, style_name: str) -> int:
        import re
        m = re.search(r"\d", style_name)
        return int(m.group()) if m else 1

    def _load_legacy_doc(self) -> List[LoadedDocument]:
        """
        FIX: The original tried to read a .doc file with PdfReader which
        is incorrect — a Word .doc is not a PDF. The correct approach is
        to use LibreOffice for conversion or textract. We now try
        LibreOffice first (converting to docx), then textract, then raise
        a clear error instead of silently returning garbage.

        """
        import subprocess, tempfile, shutil
        from pathlib import Path

        # Option 1: LibreOffice conversion
        if shutil.which("libreoffice") or shutil.which("soffice"):
            with tempfile.TemporaryDirectory() as tmpdir:
                cmd = ["libreoffice", "--headless", "--convert-to", "docx",
                       "--outdir", tmpdir, str(self.file_path)]
                try:
                    subprocess.run(cmd, capture_output=True, timeout=30, check=True)
                    converted = Path(tmpdir) / (self.file_path.stem + ".docx")
                    if converted.exists():
                        loader = DocxLoader(str(converted), config=self.config)
                        docs = loader._load_docx()
                        for doc in docs:
                            doc.metadata["source"] = str(self.file_path.resolve())
                            doc.metadata["note"] = "Converted from .doc via LibreOffice"
                        return docs
                except (subprocess.CalledProcessError, subprocess.TimeoutExpired, FileNotFoundError):
                    logger.warning("LibreOffice conversion failed, trying textract...")

        # Option 2: textract
        try:
            import textract
            text = textract.process(str(self.file_path)).decode("utf-8", errors="replace")
            meta = self._build_file_metadata(note="Loaded via textract")
            return [LoadedDocument(page_content=text.strip(), metadata=meta)]
        except ImportError:
            pass
        except Exception as e:
            logger.warning(f"textract failed: {e}")

        raise ImportError(
            f"Cannot load legacy .doc file '{self.file_path.name}'.\n"
            "Options:\n"
            "  1. Convert to .docx with Microsoft Word or LibreOffice\n"
            "  2. Install LibreOffice (used automatically)\n"
            "  3. Install textract: pip install textract"
        )
